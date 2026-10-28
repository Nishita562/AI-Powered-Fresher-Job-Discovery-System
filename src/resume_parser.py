"""In-memory resume text extraction (PDF via pypdf, DOCX via python-docx).

Uploaded files are processed from memory only: nothing is written to disk
and the text is discarded when the Streamlit session ends.
"""

from __future__ import annotations

import io

from config import settings
from src.skills import extract_skills
from src.text_utils import clean_whitespace

SUPPORTED_TYPES = {"pdf", "docx"}


class ResumeParseError(Exception):
    """Raised with a user-friendly message when a resume cannot be read."""


def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def extract_text(data: bytes, filename: str) -> str:
    ext = _ext(filename)
    if ext not in SUPPORTED_TYPES:
        raise ResumeParseError("Unsupported file type. Please upload a PDF or DOCX resume.")
    if len(data) > settings.MAX_RESUME_MB * 1024 * 1024:
        raise ResumeParseError(f"File is larger than {settings.MAX_RESUME_MB} MB.")
    if not data:
        raise ResumeParseError("The uploaded file is empty.")
    try:
        if ext == "pdf":
            if not data.startswith(b"%PDF"):
                raise ResumeParseError("This file does not look like a valid PDF.")
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ResumeParseError("The PDF is password-protected; please upload an unprotected copy.")
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        else:
            if not data.startswith(b"PK"):
                raise ResumeParseError("This file does not look like a valid DOCX document.")
            from docx import Document

            doc = Document(io.BytesIO(data))
            chunks = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    chunks.extend(cell.text for cell in row.cells)
            text = "\n".join(chunks)
    except ResumeParseError:
        raise
    except Exception as exc:  # corrupt files raise many different library errors
        raise ResumeParseError(f"Could not read the file ({type(exc).__name__}). Try another file or enter skills manually.") from exc
    text = clean_whitespace(text)
    if len(text) < 20:
        raise ResumeParseError(
            "Very little text was found — the resume may be a scanned image. Please enter your skills manually."
        )
    return text


def parse_resume(data: bytes, filename: str) -> dict:
    text = extract_text(data, filename)
    return {"text": text, "skills": extract_skills(text), "characters": len(text)}
