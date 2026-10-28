import io

import pytest
from docx import Document

from src.resume_parser import ResumeParseError, extract_text, parse_resume
from tests.conftest import minimal_pdf


def _docx_bytes(paragraphs):
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Tools: Git, Docker"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_resume_skills():
    data = _docx_bytes(["Riya Example — B.Tech AIML 2026", "Projects in Python, pandas and machine learning.", "SQL and Power BI"])
    parsed = parse_resume(data, "resume.docx")
    assert {"python", "pandas", "machine learning", "sql", "power bi", "git", "docker"} <= set(parsed["skills"])


def test_pdf_resume_text():
    data = minimal_pdf(["Arjun Example - B.Tech CSE 2026", "Skills: Python, TensorFlow, SQL, React"])
    parsed = parse_resume(data, "cv.PDF")
    assert "Arjun Example" in parsed["text"]
    assert {"python", "tensorflow", "sql", "react"} <= set(parsed["skills"])


@pytest.mark.parametrize("data, name, message", [
    (b"hello", "resume.txt", "Unsupported"),
    (b"", "resume.pdf", "empty"),
    (b"not a pdf at all", "resume.pdf", "valid PDF"),
    (b"PK\x03\x04garbage", "resume.docx", "Could not read"),
])
def test_invalid_files_raise_friendly_errors(data, name, message):
    with pytest.raises(ResumeParseError, match=message):
        extract_text(data, name)


def test_nearly_empty_resume_is_rejected():
    with pytest.raises(ResumeParseError, match="Very little text"):
        extract_text(minimal_pdf(["Hi"]), "scan.pdf")
